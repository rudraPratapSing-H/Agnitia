#!/usr/bin/env python3
"""
Build Final 12-Slide Pitch Deck (Task 3.12 - Member 4)
Generates demo/final-deck.pptx adhering to Battle Plan §13:
- 16:9 widescreen (13.333" x 7.5")
- Dark theme (#0B0F17)
- Full-sentence headlines stating the core takeaway
- One structured visual per slide
- Zero text under 24pt
- Jargon explained where it first appears
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Modern Dark Theme Colors
BG_COLOR = RGBColor(11, 15, 23)        # #0B0F17 Deep background
CARD_BG = RGBColor(22, 30, 46)         # #161E2E Elevated card
CARD_BORDER = RGBColor(51, 65, 85)     # #334155 Border slate
TEXT_PRIMARY = RGBColor(248, 250, 252) # #F8FAFC Bright white
TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Muted slate
ACCENT_CYAN = RGBColor(56, 189, 248)   # #38BDF8 Cyan
ACCENT_RED = RGBColor(248, 113, 113)   # #F87171 Red
ACCENT_AMBER = RGBColor(251, 191, 36)  # #FBBF24 Amber
ACCENT_GREEN = RGBColor(52, 211, 153)  # #34D399 Green
ACCENT_PURPLE = RGBColor(167, 139, 250)# #A78BFA Purple


def apply_dark_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def add_header(slide, tag_text, headline_text, jargon_text=None):
    # Category / context tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.size = Pt(24)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN

    # Main headline (Full sentence takeaway)
    head_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.80), Inches(11.733), Inches(0.95))
    tf_head = head_box.text_frame
    tf_head.word_wrap = True
    tf_head.margin_left = tf_head.margin_top = tf_head.margin_right = tf_head.margin_bottom = 0
    p_head = tf_head.paragraphs[0]
    p_head.text = headline_text
    p_head.font.size = Pt(32)
    p_head.font.bold = True
    p_head.font.color.rgb = TEXT_PRIMARY

    # Jargon definition / Context
    if jargon_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.80), Inches(11.733), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = jargon_text
        p_sub.font.size = Pt(24)
        p_sub.font.color.rgb = ACCENT_AMBER


def create_card(slide, left, top, width, height, border_color=CARD_BORDER, bg_color=CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(2)
    return shape


# Slide 1: Title
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • AUTONOMOUS SRE PLATFORM",
        headline_text="Agnitia: the AI on-call engineer that finds the real problem and fixes it safely.",
        jargon_text="*Jargon: SRE (Site Reliability Engineer) = software engineer who keeps production cloud services running."
    )

    create_card(slide, Inches(0.8), Inches(2.45), Inches(11.733), Inches(4.55), border_color=ACCENT_CYAN)
    tb = slide.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(11.0), Inches(3.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Built for Production Kubernetes at Hackathon 2026"
    p0.font.size = Pt(28)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN

    bullets = [
        "• Automatically correlates cascading alert storms into a single actionable incident.",
        "• Pinpoints root causes with machine-verified citations from raw cluster logs.",
        "• Synthesizes topologically ordered playbooks and awaits human authorization.",
        "• Recovers multi-tier microservice outages in under 40 seconds."
    ]
    for b in bullets:
        p = tf.add_paragraph()
        p.text = b
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY


# Slide 2: The 3 AM Problem
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="THE PROBLEM • 3 AM ON-CALL OUTAGE",
        headline_text="When one service fails, engineers get 56 alerts, not one.",
        jargon_text="*Jargon: Cascading failure = when a database outage triggers downstream connection failures in 5 other apps."
    )

    # Left: Phone Flood
    create_card(slide, Inches(0.8), Inches(2.45), Inches(5.6), Inches(4.55), border_color=ACCENT_RED)
    tb1 = slide.shapes.add_textbox(Inches(1.1), Inches(2.65), Inches(5.0), Inches(4.15))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p1 = tf1.paragraphs[0]
    p1.text = "Phone Alert Storm (56 Messages)"
    p1.font.size = Pt(26)
    p1.font.bold = True
    p1.font.color.rgb = ACCENT_RED

    alerts = [
        "• 03:14:10 [CRITICAL] Postgres: OOM Killed",
        "• 03:14:11 [ERROR] Auth: DB timeout (10x)",
        "• 03:14:11 [ERROR] Payments: DB dropped (15x)",
        "• 03:14:12 [CRITICAL] API Gateway: 503 (18x)",
        "• 03:14:14 [CRITICAL] Web UI: Checkout down (12x)"
    ]
    for a in alerts:
        p = tf1.add_paragraph()
        p.text = a
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY

    # Right: Topology chain
    create_card(slide, Inches(6.8), Inches(2.45), Inches(5.7), Inches(4.55), border_color=ACCENT_AMBER)
    tb2 = slide.shapes.add_textbox(Inches(7.1), Inches(2.65), Inches(5.1), Inches(4.15))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p2 = tf2.paragraphs[0]
    p2.text = "Cascading Dependency Chain"
    p2.font.size = Pt(26)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_AMBER

    steps = [
        "1. Database (postgres) [ROOT CAUSE: OOM]",
        "   ↓ database pool drops",
        "2. Auth & Payments [IMPACTED: 503 Errors]",
        "   ↓ upstream connection refused",
        "3. Gateway & Website [TOTAL OUTAGE]"
    ]
    for s in steps:
        p = tf2.add_paragraph()
        p.text = s
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY if not s.startswith(" ") else TEXT_MUTED


# Slide 3: Three Pains
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="THE SRE DILEMMA",
        headline_text="Finding the cause takes 30 to 60 minutes, and the wrong fix makes it worse.",
        jargon_text="*Jargon: Stampede = restarting all pods at once overwhelms the database and crashes it again."
    )

    pains = [
        ("1. THE NOISE STORM", "Engineers are flooded with 50+ notifications and waste an hour finding the root.", ACCENT_RED),
        ("2. SLOW LOG HUNTING", "Engineers manually grep megabytes of terminal logs across dozens of containers.", ACCENT_AMBER),
        ("3. INCORRECT RECOVERY", "Panicked restarts without dependency ordering cause repeated immediate crashes.", ACCENT_PURPLE)
    ]

    card_width = Inches(3.75)
    card_gap = Inches(0.24)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(pains):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.7), card_width - Inches(0.4), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


# Slide 4: The 5-Step Loop
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="THE SOLUTION • CLOSED-LOOP HEALING",
        headline_text="Agnitia closes the loop in five steps.",
        jargon_text="*Jargon: Closed loop = automated workflow from detection to verification with no manual gaps."
    )

    steps = [
        ("1. DETECT", "Streams cluster telemetry and container logs in real time.", ACCENT_CYAN),
        ("2. CORRELATE", "Collapses 56 alerts into 1 incident using graph topology.", ACCENT_PURPLE),
        ("3. DIAGNOSE", "Pinpoints root cause with verifiable log citations.", ACCENT_RED),
        ("4. PLAN", "Orders recovery steps in strict topological dependency order.", ACCENT_AMBER),
        ("5. APPROVE & HEAL", "Waits for human authorization, then heals node by node.", ACCENT_GREEN)
    ]

    card_width = Inches(2.2)
    card_gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(steps):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.65), card_width - Inches(0.3), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


# Slide 5: Live Demo Cue
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="LIVE DEMO",
        headline_text="Live demonstration: closed-loop incident resolution in action.",
        jargon_text="*Jargon: Mission control = real-time dashboard displaying topology map and streaming agent reasoning."
    )

    create_card(slide, Inches(0.8), Inches(2.45), Inches(11.733), Inches(4.55), border_color=ACCENT_GREEN)
    tb = slide.shapes.add_textbox(Inches(1.2), Inches(2.75), Inches(11.0), Inches(4.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Switching to Live Mission Control Console..."
    p0.font.size = Pt(28)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_GREEN

    items = [
        "• 1. Inject Database Out-of-Memory failure into simulated Kubernetes cluster.",
        "• 2. Watch 56 alert cards funnel into single Incident INC-104 on the topology map.",
        "• 3. Stream 5-agent pipeline investigation with verified log citation line 42.",
        "• 4. Hand judge the mobile phone to tap 'Approve' on Telegram.",
        "• 5. Observe ordered self-healing and stopwatch freeze in 38 seconds."
    ]
    for it in items:
        p = tf.add_paragraph()
        p.text = it
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY


# Slide 6: Five AI Agents & Architecture
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="SYSTEM ARCHITECTURE",
        headline_text="Five AI agents, one dependency map, and a human in control.",
        jargon_text="*Jargon: Topological order = sequence where database dependencies heal before web gateways."
    )

    agents = [
        ("1. TRIAGE", "Groups 56 raw alerts to 1 incident by evaluating root candidates.", ACCENT_CYAN),
        ("2. DIAGNOSE", "Extracts container exit code 137 and fatal memory log citations.", ACCENT_RED),
        ("3. PLAN", "Generates playbook: patch memory first, then restart in order.", ACCENT_AMBER),
        ("4. VERIFY", "Validates quotes against raw cluster logs; flags fake claims.", ACCENT_PURPLE),
        ("5. EXECUTE", "Pauses for human authorization, then heals node by node.", ACCENT_GREEN)
    ]

    card_width = Inches(2.2)
    card_gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(agents):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.65), card_width - Inches(0.3), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


# Slide 7: Proof & Verified Badges
def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="HALLUCINATION PREVENTION",
        headline_text="Every claim comes with proof the AI can't fake.",
        jargon_text="*Jargon: Exit code 137 = the system terminated the container for exceeding its memory limit."
    )

    cards = [
        ("EXIT CODE 137 [VERIFIED ✓]", [
            "• Signal: SIGKILL by Linux OOM killer",
            "• Container exceeded 64Mi memory limit",
            "• Machine-verified against cluster events"
        ], ACCENT_RED),
        ("LOG CITATION LINE 42 [VERIFIED ✓]", [
            "• Text: 'FATAL: out of memory'",
            "• Exact match confirmed in raw server logs",
            "• Fake citations drop confidence score"
        ], ACCENT_AMBER),
        ("MEMORY CURVE 64Mi [VERIFIED ✓]", [
            "• Telemetry: 40Mi ramped to 64Mi in 12s",
            "• Blast radius: 4 impacted, Redis healthy",
            "• Matches past incident INC-087 at 94%"
        ], ACCENT_PURPLE)
    ]

    card_width = Inches(3.75)
    card_gap = Inches(0.24)
    start_left = Inches(0.8)

    for i, (title, items, color) in enumerate(cards):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.65), card_width - Inches(0.4), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for it in items:
            p_i = tf.add_paragraph()
            p_i.text = it
            p_i.font.size = Pt(24)
            p_i.font.color.rgb = TEXT_PRIMARY


# Slide 8: Safety & Governance
def build_slide_8(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="SAFETY GUARDRAILS",
        headline_text="The AI can only do what we allow, and it asks before anything risky.",
        jargon_text="*Jargon: Allow-list = fixed set of 7 non-destructive actions; arbitrary shell commands are impossible."
    )

    pillars = [
        ("ALLOW-LIST (7 ACTIONS)", "Only non-destructive actions: patch memory, patch CPU, restart, rollback. No deletes.", ACCENT_CYAN),
        ("APPROVAL GATE", "Risky infrastructure changes require explicit human sign-off via UI diff or Telegram tap.", ACCENT_AMBER),
        ("AUTONOMY LEVELS 1-3", "Policy slider: Level 1 asks all, Level 2 auto-runs low risk, Level 3 stays guarded.", ACCENT_PURPLE),
        ("IMMUTABLE AUDIT LOG", "Every AI suggestion, operator authorization, and probe check is permanently logged.", ACCENT_GREEN)
    ]

    card_width = Inches(2.8)
    card_gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(pillars):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.65), card_width - Inches(0.3), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


# Slide 9: Impact & Numbers
def build_slide_9(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="PROVABLE BUSINESS METRICS",
        headline_text="In our test scenario: 56 alerts became 1 incident, and recovery took 38 seconds instead of 45 minutes.",
        jargon_text="*Jargon: MTTR (Mean Time To Recovery) = the average time required to troubleshoot and restore a down system."
    )

    metrics = [
        ("98% LESS NOISE", [
            "• 56 raw alerts collapsed to 1 incident card.",
            "• On-call engineers zero in on the root cause.",
            "• Untouched services (Redis) stay green."
        ], ACCENT_CYAN),
        ("70X FASTER RECOVERY", [
            "• Traditional manual MTTR: 30–60 minutes.",
            "• Agnitia verified recovery: exactly 38 seconds.",
            "• Prevents thousands of lost transactions."
        ], ACCENT_GREEN)
    ]

    card_width = Inches(5.7)
    card_gap = Inches(0.33)
    start_left = Inches(0.8)

    for i, (title, items, color) in enumerate(metrics):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.25), Inches(2.65), card_width - Inches(0.5), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for it in items:
            p_i = tf.add_paragraph()
            p_i.text = it
            p_i.font.size = Pt(24)
            p_i.font.color.rgb = TEXT_PRIMARY


# Slide 10: Market Comparison
def build_slide_10(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="COMPETITIVE LANDSCAPE",
        headline_text="Others group alerts or explain errors; Agnitia also fixes them in the right order.",
        jargon_text="*Jargon: Closed-loop remediation = executing the verified repair instead of merely emailing an alert."
    )

    create_card(slide, Inches(0.8), Inches(2.45), Inches(11.733), Inches(4.55), border_color=ACCENT_CYAN)
    tb = slide.shapes.add_textbox(Inches(1.1), Inches(2.7), Inches(11.1), Inches(4.1))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Capability Matrix"
    p0.font.size = Pt(26)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN

    rows = [
        "• Alert Grouping: PagerDuty & BigPanda group alerts; Agnitia uses the live graph.",
        "• Machine Evidence: Enterprise tools partly check logs; Agnitia verifies citations.",
        "• Ordered Execution: Others require manual scripts; Agnitia builds topological orders.",
        "• Infrastructure Overhead: Big platforms cost $$$/host; Agnitia reads K8s directly.",
        "• Early Warning: Agnitia predicts exhaustion countdown before the OOM crash happens."
    ]
    for r in rows:
        p = tf.add_paragraph()
        p.text = r
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY


# Slide 11: Business & Who Pays
def build_slide_11(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="TARGET MARKET & BUSINESS VALUE",
        headline_text="Built for teams running Kubernetes without a 24/7 reliability team.",
        jargon_text="*Jargon: 24/7 on-call rotation = team members waking up in shifts to answer midnight production outages."
    )

    sections = [
        ("WHO PAYS FOR AGNITIA", [
            "• Growth-stage startups running 10–50 microservices.",
            "• Engineering teams without dedicated 24/7 SRE staff.",
            "• Saves ~₹1,00,000+ per midnight downtime incident."
        ], ACCENT_AMBER),
        ("ROADMAP TRAJECTORY", [
            "• Auto-discovered topologies via eBPF / OpenTelemetry.",
            "• Native Slack, PagerDuty, and Opsgenie push integrations.",
            "• Multi-cluster Kubernetes and cloud mesh support."
        ], ACCENT_PURPLE)
    ]

    card_width = Inches(5.7)
    card_gap = Inches(0.33)
    start_left = Inches(0.8)

    for i, (title, items, color) in enumerate(sections):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.25), Inches(2.65), card_width - Inches(0.5), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(28)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for it in items:
            p_i = tf.add_paragraph()
            p_i.text = it
            p_i.font.size = Pt(24)
            p_i.font.color.rgb = TEXT_PRIMARY


# Slide 12: Team & Conclusion
def build_slide_12(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="CONCLUSION & WRAP-UP",
        headline_text="Built in 18 hours: from 56 alerts to one verified fix, safely.",
        jargon_text="*Jargon: Closed-loop reliability = AI finding, proving, and healing infrastructure outages."
    )

    create_card(slide, Inches(0.8), Inches(2.45), Inches(11.733), Inches(4.55), border_color=ACCENT_GREEN)
    tb = slide.shapes.add_textbox(Inches(1.2), Inches(2.8), Inches(11.0), Inches(3.9))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "Team Agnitia • 4 Engineers, 4 Lanes, 1 Shared Contract"
    p0.font.size = Pt(28)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_GREEN

    team = [
        "• Member 1: Frontend Lead (React Flow topology map, Evidence drawer, Chaos bar)",
        "• Member 2: Backend & Simulator Lead (Deterministic cluster simulator, Executor, Adapters)",
        "• Member 3: AI Agents Lead (Diagnosis, Planner, Citation verifier, Postmortem generator)",
        "• Member 4: Chaos Content, Pitch & Integrations (Scenarios, Telegram bot, Decks, Audio)",
        "• Thank you! Agnitia is ready for your questions."
    ]
    for m in team:
        p = tf.add_paragraph()
        p.text = m
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY


def ensure_min_font_size(prs, min_pt=24):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if not p.font.size or p.font.size.pt < min_pt:
                        p.font.size = Pt(min_pt)
                    for run in p.runs:
                        if not run.font.size or run.font.size.pt < min_pt:
                            run.font.size = Pt(min_pt)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)
    build_slide_8(prs)
    build_slide_9(prs)
    build_slide_10(prs)
    build_slide_11(prs)
    build_slide_12(prs)

    ensure_min_font_size(prs, min_pt=24)

    repo_root = Path(__file__).resolve().parent.parent
    demo_dir = repo_root / "demo"
    final_dir = demo_dir / "final-deck"
    final_dir.mkdir(parents=True, exist_ok=True)

    out1 = demo_dir / "final-deck.pptx"
    out2 = final_dir / "final-deck.pptx"

    prs.save(str(out1))
    prs.save(str(out2))

    print("Saved Final 12-slide deck successfully to:")
    print(f" - {out1.name}")
    print(f" - final-deck/{out2.name}")


if __name__ == "__main__":
    main()
