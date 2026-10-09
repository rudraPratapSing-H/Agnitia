#!/usr/bin/env python3
"""
Build Checkpoint 1 PowerPoint Deck (Task 1.13 - Member 4)
Generates demo/cp1-deck.pptx with 16:9 widescreen, dark theme, full-sentence headlines,
one visual per slide, no text under 24pt, and jargon explained where it first appears.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE


# 16:9 Widescreen dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette (Dark Theme)
BG_COLOR = RGBColor(11, 15, 23)        # #0B0F17 Dark background
CARD_BG = RGBColor(22, 30, 46)         # #161E2E Card background
CARD_BORDER = RGBColor(51, 65, 85)     # #334155 Border
TEXT_PRIMARY = RGBColor(248, 250, 252) # #F8FAFC White
TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Muted slate
ACCENT_CYAN = RGBColor(56, 189, 248)   # #38BDF8 Cyan
ACCENT_RED = RGBColor(248, 113, 113)   # #F87171 Red
ACCENT_AMBER = RGBColor(251, 191, 36)  # #FBBF24 Amber
ACCENT_GREEN = RGBColor(52, 211, 153)  # #34D399 Green
ACCENT_PURPLE = RGBColor(167, 139, 250)# #A78BFA Purple


def apply_dark_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def add_header(slide, tag_text, headline_text, jargon_text=None):
    # Category tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.size = Pt(24)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN

    # Main headline (Full sentence takeaway)
    head_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.733), Inches(0.85))
    tf_head = head_box.text_frame
    tf_head.word_wrap = True
    tf_head.margin_left = tf_head.margin_top = tf_head.margin_right = tf_head.margin_bottom = 0
    p_head = tf_head.paragraphs[0]
    p_head.text = headline_text
    p_head.font.size = Pt(32)
    p_head.font.bold = True
    p_head.font.color.rgb = TEXT_PRIMARY

    # Jargon definition / Context
    if jargon_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.75), Inches(11.733), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = jargon_text
        p_sub.font.size = Pt(24)
        p_sub.font.color.rgb = ACCENT_AMBER


def create_card(slide, left, top, width, height, border_color=CARD_BORDER, bg_color=CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(2)
    return shape


def build_slide_1(prs):
    # Slide 1: The 3 AM Problem
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • CHECKPOINT 1 • THE PROBLEM",
        headline_text="The 3 AM problem: when one service fails, engineers get 56 alerts, not one.",
        jargon_text="*Jargon: 56 alerts = cascade of error messages sent to phones when downstream systems fail."
    )

    # Visual: Left card (Phone Alert Flood) vs Right card (Failure Chain)
    # Left Card - Flooded Phone
    create_card(slide, Inches(0.8), Inches(2.4), Inches(5.6), Inches(4.5), border_color=ACCENT_RED)
    tb_phone = slide.shapes.add_textbox(Inches(1.1), Inches(2.6), Inches(5.0), Inches(4.1))
    tf_p = tb_phone.text_frame
    tf_p.word_wrap = True

    p0 = tf_p.paragraphs[0]
    p0.text = "Phone Alert Storm (56 Messages)"
    p0.font.size = Pt(26)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_RED

    alerts = [
        "03:14:10 [CRITICAL] Database OOM Killed (code 137)",
        "03:14:11 [ERROR] Auth: DB query connection timeout",
        "03:14:11 [ERROR] Payments: DB connection refused",
        "03:14:12 [CRITICAL] API Gateway: 502/503 spike (18x)",
        "03:14:14 [CRITICAL] Web UI: Checkout checkout failed"
    ]
    for text in alerts:
        p = tf_p.add_paragraph()
        p.text = f"• {text}"
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY

    # Right Card - Cascading Dependency Chain
    create_card(slide, Inches(6.8), Inches(2.4), Inches(5.7), Inches(4.5), border_color=ACCENT_AMBER)
    tb_chain = slide.shapes.add_textbox(Inches(7.1), Inches(2.6), Inches(5.1), Inches(4.1))
    tf_c = tb_chain.text_frame
    tf_c.word_wrap = True

    p_c0 = tf_c.paragraphs[0]
    p_c0.text = "Cascading Failure Chain"
    p_c0.font.size = Pt(26)
    p_c0.font.bold = True
    p_c0.font.color.rgb = ACCENT_AMBER

    steps = [
        "1. Database (postgres) [ROOT CAUSE: OOM]",
        "   ↓ failure propagates to dependencies",
        "2. Auth & Payments [IMPACTED: 503 Errors]",
        "   ↓ gateway upstream connections fail",
        "3. Gateway & Website [FULL SYSTEM OUTAGE]"
    ]
    for s in steps:
        p = tf_c.add_paragraph()
        p.text = s
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY if not s.startswith(" ") else TEXT_MUTED


def build_slide_2(prs):
    # Slide 2: Five Steps
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • THE SOLUTION LOOP",
        headline_text="Agnitia closes the loop in five steps.",
        jargon_text="*Jargon: Closed loop = automated workflow from incident detection to verified resolution."
    )

    # Visual: 5 sequential cards horizontally arranged
    step_data = [
        ("1. DETECT", "Streams cluster metrics and container logs in real time.", ACCENT_CYAN),
        ("2. CORRELATE", "Collapses 56 alerts into 1 incident using topology graph.", ACCENT_PURPLE),
        ("3. DIAGNOSE", "Pinpoints root cause with verifiable log citations.", ACCENT_RED),
        ("4. PLAN", "Generates safe, dependency-ordered recovery actions.", ACCENT_AMBER),
        ("5. APPROVE & HEAL", "Waits for human authorization, then verifies recovery.", ACCENT_GREEN)
    ]

    card_width = Inches(2.2)
    card_gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(step_data):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.5), card_width, Inches(4.4), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.7), card_width - Inches(0.3), Inches(4.0))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


def build_slide_3(prs):
    # Slide 3: Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • SYSTEM ARCHITECTURE",
        headline_text="One backend, a live dependency map and a swappable cluster.",
        jargon_text="*Jargon: Cluster adapter = interface switching between test simulator and real Kubernetes."
    )

    # Visual: 3 horizontal layered tiers
    layers = [
        ("TIER 1 • REACT OPERATOR UI", "Live React Flow dependency map • Terminal AI reasoning panel • Incident approval card", ACCENT_CYAN),
        ("TIER 2 • FASTAPI & EVENT BUS", "Topology graph correlation • Real-time WebSocket bus • Pydantic contract validation", ACCENT_PURPLE),
        ("TIER 3 • CLUSTER ADAPTERS", "Deterministic simulator (demo mode) OR Real Kubernetes cluster (Kind)", ACCENT_GREEN)
    ]

    top_pos = Inches(2.5)
    layer_height = Inches(1.3)
    layer_gap = Inches(0.25)

    for i, (title, desc, color) in enumerate(layers):
        curr_top = top_pos + i * (layer_height + layer_gap)
        create_card(slide, Inches(0.8), curr_top, Inches(11.733), layer_height, border_color=color)

        tb = slide.shapes.add_textbox(Inches(1.1), curr_top + Inches(0.15), Inches(11.1), layer_height - Inches(0.3))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


def build_slide_4(prs):
    # Slide 4: Progress Board
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • EXECUTION VELOCITY",
        headline_text="Where we are: progress board.",
        jargon_text="*Jargon: Chaos panel = one-click fault injection tool to verify resilience under fire."
    )

    # Visual: 3 Columns (Done, Building now, Next) using feature names from battle plan §3
    cols = [
        ("DONE (HOUR 5)", [
            "• Live Dependency Map",
            "• Chaos Injection Panel",
            "• Alert Storm Funnel",
            "• Live AI Reasoning Panel",
            "• Telegram Bot Ping/Pong"
        ], ACCENT_GREEN),
        ("BUILDING NOW", [
            "• Evidence Drawer with Proof",
            "• Ordered Playbook Engine",
            "• Live 5-Agent Pipeline",
            "• Citation Verifier Guard"
        ], ACCENT_AMBER),
        ("NEXT (CHECKPOINT 2)", [
            "• Approval Card & Diff",
            "• Self-Healing Animation",
            "• Mobile 1-Click Approval",
            "• Recovery Cost Stopwatch"
        ], ACCENT_CYAN)
    ]

    col_width = Inches(3.75)
    col_gap = Inches(0.24)
    start_left = Inches(0.8)

    for i, (title, items, color) in enumerate(cols):
        left = start_left + i * (col_width + col_gap)
        create_card(slide, left, Inches(2.5), col_width, Inches(4.5), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.7), col_width - Inches(0.4), Inches(4.1))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = item
            p_i.font.size = Pt(24)
            p_i.font.color.rgb = TEXT_PRIMARY


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)

    repo_root = Path(__file__).resolve().parent.parent
    demo_dir = repo_root / "demo"
    deck_dir = demo_dir / "cp1-deck"
    deck_dir.mkdir(parents=True, exist_ok=True)

    # Save to demo/cp1-deck.pptx and demo/cp1-deck/cp1-deck.pptx
    out1 = demo_dir / "cp1-deck.pptx"
    out2 = deck_dir / "cp1-deck.pptx"

    prs.save(str(out1))
    prs.save(str(out2))

    print(f"Saved Checkpoint 1 deck successfully to:")
    print(f" - {out1.name}")
    print(f" - cp1-deck/{out2.name}")


if __name__ == "__main__":
    main()
