#!/usr/bin/env python3
"""
Build Checkpoint 2 PowerPoint Deck (Task 2.16 - Member 4)
Generates demo/cp2-deck.pptx with 16:9 widescreen, dark theme, full-sentence headlines,
one visual per slide, no text under 24pt, and jargon explained where it first appears.
Covers Battle Plan Slides 6-8 + Checkpoint 2 Progress Board & Promise.
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


# 16:9 Widescreen dimensions
SLIDE_WIDTH = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)

# Color Palette (Dark Theme)
BG_COLOR = RGBColor(11, 15, 23)        # #0B0F17 Dark background
CARD_BG = RGBColor(22, 30, 46)         # #161E2E Card background
CARD_BORDER = RGBColor(51, 65, 85)     # #334155 Border
TEXT_PRIMARY = RGBColor(248, 250, 252) # #F8FAFC White
TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Muted slate
ACCENT_CYAN = RGBColor(56, 189, 248)   # #38BDF8 Cyan
ACCENT_RED = RGBColor(248, 113, 113)   # #F87171 Red
ACCENT_AMBER = RGBColor(251, 191, 36)  # #FBBF24 Amber
ACCENT_GREEN = RGBColor(52, 211, 153)  # #34D399 Green
ACCENT_PURPLE = RGBColor(167, 139, 250)# #A78BFA Purple


def apply_dark_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def add_header(slide, tag_text, headline_text, jargon_text=None):
    # Category tag
    tag_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(0.4))
    tf_tag = tag_box.text_frame
    tf_tag.word_wrap = True
    tf_tag.margin_left = tf_tag.margin_top = tf_tag.margin_right = tf_tag.margin_bottom = 0
    p_tag = tf_tag.paragraphs[0]
    p_tag.text = tag_text.upper()
    p_tag.font.size = Pt(24)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN

    # Main headline (Full sentence takeaway)
    head_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.80), Inches(11.733), Inches(0.95))
    tf_head = head_box.text_frame
    tf_head.word_wrap = True
    tf_head.margin_left = tf_head.margin_top = tf_head.margin_right = tf_head.margin_bottom = 0
    p_head = tf_head.paragraphs[0]
    p_head.text = headline_text
    p_head.font.size = Pt(32)
    p_head.font.bold = True
    p_head.font.color.rgb = TEXT_PRIMARY

    # Jargon definition / Context
    if jargon_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.80), Inches(11.733), Inches(0.45))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        tf_sub.margin_left = tf_sub.margin_top = tf_sub.margin_right = tf_sub.margin_bottom = 0
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = jargon_text
        p_sub.font.size = Pt(24)
        p_sub.font.color.rgb = ACCENT_AMBER


def create_card(slide, left, top, width, height, border_color=CARD_BORDER, bg_color=CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(2)
    return shape


# ----------------------------------------------------------------------
# Slide 1: Recap - The 3 AM Problem
# ----------------------------------------------------------------------
def build_slide_1(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • RECAP • THE PROBLEM",
        headline_text="The 3 AM problem: when one service fails, engineers get 56 alerts, not one.",
        jargon_text="*Jargon: Alert storm = cascading notifications flooding phones when downstream services fail."
    )

    # Left Card - Flooded Phone
    create_card(slide, Inches(0.8), Inches(2.45), Inches(5.6), Inches(4.55), border_color=ACCENT_RED)
    tb_phone = slide.shapes.add_textbox(Inches(1.1), Inches(2.65), Inches(5.0), Inches(4.15))
    tf_p = tb_phone.text_frame
    tf_p.word_wrap = True

    p0 = tf_p.paragraphs[0]
    p0.text = "Phone Alert Storm (56 Messages)"
    p0.font.size = Pt(26)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_RED

    alerts = [
        "03:14:10 [CRITICAL] Database OOM Killed (code 137)",
        "03:14:11 [ERROR] Auth: DB connection timeout",
        "03:14:11 [ERROR] Payments: DB connection refused",
        "03:14:12 [CRITICAL] API Gateway: 503 spike (18x)",
        "03:14:14 [CRITICAL] Web UI: Checkout failed (12x)"
    ]
    for text in alerts:
        p = tf_p.add_paragraph()
        p.text = f"• {text}"
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY

    # Right Card - Cascading Dependency Chain
    create_card(slide, Inches(6.8), Inches(2.45), Inches(5.7), Inches(4.55), border_color=ACCENT_AMBER)
    tb_chain = slide.shapes.add_textbox(Inches(7.1), Inches(2.65), Inches(5.1), Inches(4.15))
    tf_c = tb_chain.text_frame
    tf_c.word_wrap = True

    p_c0 = tf_c.paragraphs[0]
    p_c0.text = "Cascading Failure Topology"
    p_c0.font.size = Pt(26)
    p_c0.font.bold = True
    p_c0.font.color.rgb = ACCENT_AMBER

    steps = [
        "1. Database (postgres) [ROOT CAUSE: OOM]",
        "   ↓ failure cascades to dependent services",
        "2. Auth & Payments [IMPACTED: 503 Outage]",
        "   ↓ gateway upstream pools exhaust",
        "3. Gateway & Website [FULL USER OUTAGE]"
    ]
    for s in steps:
        p = tf_c.add_paragraph()
        p.text = s
        p.font.size = Pt(24)
        p.font.color.rgb = TEXT_PRIMARY if not s.startswith(" ") else TEXT_MUTED


# ----------------------------------------------------------------------
# Slide 2: The Solution Loop
# ----------------------------------------------------------------------
def build_slide_2(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • CLOSED-LOOP AUTOMATION",
        headline_text="Agnitia closes the loop in five steps.",
        jargon_text="*Jargon: Closed loop = automated progression from detection to verified recovery without manual delay."
    )

    step_data = [
        ("1. DETECT", "Streams cluster metrics and container logs in real time.", ACCENT_CYAN),
        ("2. CORRELATE", "Collapses 56 alerts into 1 incident using graph topology.", ACCENT_PURPLE),
        ("3. DIAGNOSE", "Pinpoints root cause with verifiable log citations.", ACCENT_RED),
        ("4. PLAN", "Orders recovery steps in strict topological dependency order.", ACCENT_AMBER),
        ("5. APPROVE & HEAL", "Waits for human authorization, then heals node by node.", ACCENT_GREEN)
    ]

    card_width = Inches(2.2)
    card_gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(step_data):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.65), card_width - Inches(0.3), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


# ----------------------------------------------------------------------
# Slide 3 (Battle Plan Slide 6): Five AI Agents & Architecture
# ----------------------------------------------------------------------
def build_slide_3(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • CHECKPOINT 2 • AGENT PIPELINE",
        headline_text="Five AI agents, one dependency map, and a human in control.",
        jargon_text="*Jargon: Topological order = recovery sequence where foundational databases heal before web services."
    )

    agents = [
        ("1. TRIAGE", "Groups 56 raw alerts into 1 incident by checking graph roots.", ACCENT_CYAN),
        ("2. DIAGNOSE", "Extracts container exit code 137 and fatal memory log citations.", ACCENT_RED),
        ("3. PLAN", "Generates playbook: patch memory first, then restart in order.", ACCENT_AMBER),
        ("4. VERIFY", "Validates all log quotes against raw data; flags fake claims.", ACCENT_PURPLE),
        ("5. EXECUTE", "Pauses for human approval, then verifies health per node.", ACCENT_GREEN)
    ]

    card_width = Inches(2.2)
    card_gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(agents):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.65), card_width - Inches(0.3), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


# ----------------------------------------------------------------------
# Slide 4 (Battle Plan Slide 7): Machine-Checked Citations (Evidence Drawer)
# ----------------------------------------------------------------------
def build_slide_4(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • CHECKPOINT 2 • EVIDENCE & PROOF",
        headline_text="Every claim comes with proof the AI can't fake.",
        jargon_text="*Jargon: Exit code 137 = the Linux kernel killed the container for exceeding its memory limit."
    )

    # 3 Evidence verification cards
    cards_data = [
        ("EXIT CODE 137 [VERIFIED ✓]", [
            "• Signal: SIGKILL by Linux OOM killer",
            "• Reason: Container exceeded 64Mi limit",
            "• Checked: Machine-verified via K8s events"
        ], ACCENT_RED),
        ("LOG CITATION LINE 42 [VERIFIED ✓]", [
            "• Text: 'FATAL: out of memory'",
            "• Exact Match: Found at line 42 in raw log",
            "• Guard: Hallucinated citations dropped"
        ], ACCENT_AMBER),
        ("MEMORY CURVE 64Mi [VERIFIED ✓]", [
            "• Telemetry: 40Mi ramped to 64Mi in 12s",
            "• Blast Radius: 4 victims, Redis healthy",
            "• Incident Match: 94% match to INC-087"
        ], ACCENT_PURPLE)
    ]

    card_width = Inches(3.75)
    card_gap = Inches(0.24)
    start_left = Inches(0.8)

    for i, (title, items, color) in enumerate(cards_data):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.65), card_width - Inches(0.4), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = item
            p_i.font.size = Pt(24)
            p_i.font.color.rgb = TEXT_PRIMARY


# ----------------------------------------------------------------------
# Slide 5 (Battle Plan Slide 8): Safety & Governance
# ----------------------------------------------------------------------
def build_slide_5(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • CHECKPOINT 2 • SAFETY & GOVERNANCE",
        headline_text="The AI can only do what we allow, and it asks before anything risky.",
        jargon_text="*Jargon: Allow-list = fixed set of 7 non-destructive actions; arbitrary shell commands are forbidden."
    )

    safety_pillars = [
        ("ALLOW-LIST (7 ACTIONS)", "Only permitted operations: patch memory, patch CPU, restart, rollback. Deletions impossible.", ACCENT_CYAN),
        ("APPROVAL GATE", "Risky config changes require human authorization via UI diff or mobile Telegram tap.", ACCENT_AMBER),
        ("AUTONOMY LEVELS 1-3", "Policy slider: Level 1 asks all, Level 2 auto-runs low risk, Level 3 stays guarded.", ACCENT_PURPLE),
        ("IMMUTABLE AUDIT LOG", "Every AI suggestion, human approval, and probe check is permanently recorded.", ACCENT_GREEN)
    ]

    card_width = Inches(2.8)
    card_gap = Inches(0.18)
    start_left = Inches(0.8)

    for i, (title, desc, color) in enumerate(safety_pillars):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.15), Inches(2.65), card_width - Inches(0.3), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(24)
        p_d.font.color.rgb = TEXT_PRIMARY


# ----------------------------------------------------------------------
# Slide 6: Checkpoint 2 Progress Board
# ----------------------------------------------------------------------
def build_slide_6(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • EXECUTION VELOCITY",
        headline_text="Where we are: Checkpoint 2 progress board.",
        jargon_text="*Jargon: Feature freeze = hour 15 milestone when new development locks to ensure stable rehearsal."
    )

    cols = [
        ("DONE (CHECKPOINT 2)", [
            "• Live 5-Agent Pipeline",
            "• Evidence Drawer with Citations",
            "• Ordered Playbook & Diff",
            "• Mobile Approval Bot (Telegram)",
            "• Downtime Cost Stopwatch"
        ], ACCENT_GREEN),
        ("BUILDING NOW (PHASE 3)", [
            "• Crash Prediction (Slow Leak)",
            "• Real-Kubernetes Adapter (Kind)",
            "• Autonomy Policy Slider (L1-3)",
            "• Instant Postmortem Writer",
            "• Past Incidents Memory (INC-087)"
        ], ACCENT_AMBER),
        ("NEXT (FINAL REHEARSAL)", [
            "• 10 Timed Stage Rehearsals",
            "• Projector 1080p Calibration",
            "• Demo Safety Net Freeze",
            "• Q&A Team Drill (§14)"
        ], ACCENT_CYAN)
    ]

    col_width = Inches(3.75)
    col_gap = Inches(0.24)
    start_left = Inches(0.8)

    for i, (title, items, color) in enumerate(cols):
        left = start_left + i * (col_width + col_gap)
        create_card(slide, left, Inches(2.45), col_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.2), Inches(2.65), col_width - Inches(0.4), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = item
            p_i.font.size = Pt(24)
            p_i.font.color.rgb = TEXT_PRIMARY


# ----------------------------------------------------------------------
# Slide 7: Checkpoint 2 Closing Promise
# ----------------------------------------------------------------------
def build_slide_7(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_dark_background(slide)
    add_header(
        slide,
        tag_text="AGNITIA • ROAD TO FINAL",
        headline_text="In the final, Agnitia will predict crashes before they happen.",
        jargon_text="*Jargon: Preventive fix = raising resources before an outage occurs to avoid any user downtime."
    )

    features = [
        ("FEATURE 16 • PREVENTIVE CRASH PREDICTION", [
            "• Linear regression detects memory climbs 4 minutes early.",
            "• Countdown warning fires at t=180s with 60s safety buffer.",
            "• Offers preventive memory patch before the OOM crash happens.",
            "• Proves proactive reliability on the slow_leak scenario."
        ], ACCENT_CYAN),
        ("FEATURE 14 • AUTONOMOUS POSTMORTEM WRITER", [
            "• Generates comprehensive blameless report in 10 seconds.",
            "• Includes verified timeline, root cause evidence, and impact.",
            "• Outputs 3 role-assigned action items (Platform, SRE, Dev).",
            "• Fully downloadable markdown document with one click."
        ], ACCENT_PURPLE)
    ]

    card_width = Inches(5.7)
    card_gap = Inches(0.33)
    start_left = Inches(0.8)

    for i, (title, items, color) in enumerate(features):
        left = start_left + i * (card_width + card_gap)
        create_card(slide, left, Inches(2.45), card_width, Inches(4.55), border_color=color)

        tb = slide.shapes.add_textbox(left + Inches(0.25), Inches(2.65), card_width - Inches(0.5), Inches(4.15))
        tf = tb.text_frame
        tf.word_wrap = True

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(26)
        p_t.font.bold = True
        p_t.font.color.rgb = color

        for item in items:
            p_i = tf.add_paragraph()
            p_i.text = item
            p_i.font.size = Pt(24)
            p_i.font.color.rgb = TEXT_PRIMARY


def ensure_min_font_size(prs, min_pt=24):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if not p.font.size or p.font.size.pt < min_pt:
                        p.font.size = Pt(min_pt)
                    for run in p.runs:
                        if not run.font.size or run.font.size.pt < min_pt:
                            run.font.size = Pt(min_pt)


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT

    build_slide_1(prs)
    build_slide_2(prs)
    build_slide_3(prs)
    build_slide_4(prs)
    build_slide_5(prs)
    build_slide_6(prs)
    build_slide_7(prs)

    ensure_min_font_size(prs, min_pt=24)

    repo_root = Path(__file__).resolve().parent.parent
    demo_dir = repo_root / "demo"
    deck_dir = demo_dir / "cp2-deck"
    deck_dir.mkdir(parents=True, exist_ok=True)

    out1 = demo_dir / "cp2-deck.pptx"
    out2 = deck_dir / "cp2-deck.pptx"

    prs.save(str(out1))
    prs.save(str(out2))

    print("Saved Checkpoint 2 deck successfully to:")
    print(f" - {out1.name}")
    print(f" - cp2-deck/{out2.name}")


if __name__ == "__main__":
    main()
